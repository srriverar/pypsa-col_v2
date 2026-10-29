"""Import official PEN 2022–2052 LEAP *results* (anexo UPME).

The .leap files are NexusDB (encrypted in 2022). We cannot execute LEAP here.
The public anexo is the authorised numeric export of those runs.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

SCENARIO_SHEETS = {
    "A": "Actualización",
    "M": "Modernización",
    "I": "Inflexión",
    "D": "Disrupción",
    "LS": "Transición límite superior",
    "LI": "Transición límite inferior",
}

SECTOR_START = {
    "transporte": 3,
    "industria": 22,
    "residencial": 40,
    "terciario": 58,
    "agro_const_mineria": 76,
    "no_energetico": 94,
}

SHEET_KIND = {
    "Consumo Final": "consumo_final_pj",
    "Emisiones Consumo Final": "emisiones_final_mtco2e",
}


def _root() -> Path:
    return Path(__file__).resolve().parents[3]


def anexo_path() -> Path:
    return _root() / "data" / "mde" / "leap" / "Anexo_Resultados_PEN_2022_2052.xlsx"


def leap_files() -> dict[str, Path]:
    d = _root() / "data" / "mde" / "leap"
    return {
        "pen_2022_leap": d / "escbase2022v26_20_06_2023.leap",
        "pen_2020_leap": d / "escbase2020v25_13_11_2020.leap",
        "anexo_xlsx": anexo_path(),
        "supuestos_2020": d / "Supuestos_PEN_2020_2050.xlsx",
    }


def _parse_block(ws, start_col: int, header_row: int = 11) -> tuple[list[str], int]:
    fuels = []
    col = start_col
    while True:
        v = ws.cell(header_row, col).value
        if v is None:
            break
        if str(v).strip().lower() in {"año", "ano", "year"}:
            break
        fuels.append(str(v).strip())
        col += 1
    return fuels, col


def parse_anexo(path: Path | None = None) -> pd.DataFrame:
    import openpyxl

    path = path or anexo_path()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    rows: list[dict] = []
    for code, scenario in SCENARIO_SHEETS.items():
        for prefix, kind in SHEET_KIND.items():
            name = f"{prefix} {code}"
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            for sector, start in SECTOR_START.items():
                fuels, _end = _parse_block(ws, start)
                if not fuels:
                    continue
                year_col = start - 1
                for r in range(12, ws.max_row + 1):
                    year = ws.cell(r, year_col).value
                    if year is None:
                        continue
                    try:
                        year_i = int(year)
                    except (TypeError, ValueError):
                        continue
                    if year_i < 1990 or year_i > 2100:
                        continue
                    for i, fuel in enumerate(fuels):
                        val = ws.cell(r, start + i).value
                        if val is None:
                            continue
                        try:
                            num = float(val)
                        except (TypeError, ValueError):
                            continue
                        rows.append(
                            {
                                "source": "PEN_2022_2052_anexo_LEAP",
                                "scenario_code": code,
                                "scenario": scenario,
                                "kind": kind,
                                "sector": sector,
                                "fuel": fuel,
                                "year": year_i,
                                "value": num,
                            }
                        )
    wb.close()
    return pd.DataFrame(rows)


def load_pen_results(force: bool = False) -> pd.DataFrame:
    cache = _root() / "data" / "mde" / "pen_consumo_emisiones.csv"
    if cache.exists() and not force:
        return pd.read_csv(cache)
    df = parse_anexo()
    cache.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(cache, index=False)
    return df


if __name__ == "__main__":
    df = load_pen_results(force=True)
    print(df.shape, df.scenario.unique().tolist(), df.sector.unique().tolist())
    print(df.groupby(["kind", "scenario"]).value.sum().head(20))
